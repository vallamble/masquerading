# -*- coding: utf-8 -*-
"""
make_gap_dataset.py — Dataset d'ÉCART (100 % fictif) pour les exigences Novartis non couvertes :
  A. RTF                      -> Lettre_FICTIF.rtf (DOCX python-docx -> LibreOffice -> RTF)
  B. PDF scanné               -> Rapport_scan_FICTIF.pdf (4 pages image 200 dpi, skew, bruit, JPEG q75, une page /Rotate 90)
  C. documents imbriqués      -> Contrat_imbrique_FICTIF.docx (XLSX en objet OLE word/embeddings + aperçu PNG en clair)
                                 Annexe_jointe_FICTIF.pdf (pièce jointe XLSX via embfile_add)
  D. signatures manuscrites   -> Formulaire_signe_FICTIF.pdf (p.1 paraphe raster, p.2 paraphe vectoriel) + formulaire_signe_FICTIF.png

Écrit UNIQUEMENT dans <DS>/inbound_gap/ et <DS>/ground_truth/{ground_truth_gap,signatures_gap,decoys_gap}.csv.
Les valeurs viennent de ground_truth.csv / mapping_by_value.csv (mêmes pseudonymes que le reste du dataset).

    python tools/make_gap_dataset.py --ds <…/02_Phase2_dataset> [--soffice /Applications/LibreOffice.app/Contents/MacOS/soffice]
"""
import argparse
import collections
import csv
import io
import math
import os
import random
import shutil
import subprocess
import sys
import tempfile
import zipfile

import fitz  # PyMuPDF
import numpy as np
from PIL import Image, ImageDraw, ImageFilter, ImageFont

SEED = 20260916
GT_COLS = ["gt_id", "document", "location_type", "page", "sheet", "cell", "paragraph", "image_file", "section",
           "data_type", "value", "entity_id", "entity_role", "variant", "language", "expected_action",
           "expected_replacement", "strict_replacement", "securiti_data_element_hint", "native_or_custom"]
HINT = {"FIRST_NAME": "First Name", "LAST_NAME": "Last Name", "DATE_OF_BIRTH": "Date of Birth", "AHV_NUMBER": "CH AHV",
        "IBAN": "IBAN (Global)", "PATIENT_ID": "custom PAT-\\d{4}-\\d{5}", "INSURANCE_CARD_NUMBER": "Health Insurance Plan Member ID",
        "STREET_ADDRESS": "Street Name", "POSTAL_CITY": "Postal Code + City Name", "PHONE": "Phone Number", "EMAIL": "Email Address"}
SIGN_LABEL_FONT = "/System/Library/Fonts/Supplemental/Arial.ttf"
FONT_CANDIDATES = [SIGN_LABEL_FONT, "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/Library/Fonts/Arial.ttf"]


def font(size):
    for p in FONT_CANDIDATES:
        if os.path.exists(p):
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()


# ---------------------------------------------------------------------------------------------------------------------
class GT:
    """Accumule les lignes de vérité terrain, de leurres et de boîtes de signature."""

    def __init__(self, entities):
        self.rows, self.decoys, self.sigs = [], [], []
        self.ents = entities
        self.n = 0

    def add(self, document, location_type, dtype, value, entity_id, *, page="", sheet="", cell="", paragraph="",
            image_file="", section="", variant="text", language="fr", role="patient"):
        rep = self.ents[entity_id][dtype][1]
        self.n += 1
        self.rows.append(dict(gt_id="GAP%05d" % self.n, document=document, location_type=location_type, page=page, sheet=sheet,
                              cell=cell, paragraph=paragraph, image_file=image_file, section=section, data_type=dtype, value=value,
                              entity_id=entity_id, entity_role=role, variant=variant, language=language, expected_action="PSEUDONYMIZE",
                              expected_replacement=rep, strict_replacement=rep, securiti_data_element_hint=HINT.get(dtype, dtype),
                              native_or_custom="custom" if dtype in ("PATIENT_ID", "INSURANCE_CARD_NUMBER") else "natif"))

    def add_person(self, document, location_type, eid, types, **kw):
        for t in types:
            self.add(document, location_type, t, self.ents[eid][t][0], eid, **kw)

    def decoy(self, document, kind, value, note, page="", sheet="", cell=""):
        self.decoys.append(dict(decoy_id="DKG%03d" % (len(self.decoys) + 1), document=document, page=page, section="", sheet=sheet,
                                cell=cell, kind=kind, value=value, expected="NOT detected (faux positif si détecté)", note=note))

    def sig(self, file, page, box, kind, unit):
        self.sigs.append(dict(file=file, page=page, x0=round(box[0], 1), y0=round(box[1], 1), x1=round(box[2], 1), y1=round(box[3], 1),
                              kind=kind, unit=unit))


def load_entities(ds):
    rows = list(csv.DictReader(open(os.path.join(ds, "ground_truth/ground_truth.csv"), encoding="utf-8")))
    ents = collections.defaultdict(dict)
    for r in rows:
        if r["entity_role"] == "patient":
            ents[r["entity_id"]].setdefault(r["data_type"], (r["value"], r["expected_replacement"]))
    return rows, ents


def full_name(ents, eid):
    return "%s %s" % (ents[eid]["FIRST_NAME"][0], ents[eid]["LAST_NAME"][0])


# ---------------------------------------------------------------------------------------------------------------------
# B. PDF scanné
# ---------------------------------------------------------------------------------------------------------------------
def degrade(img, rng, angle):
    """Rotation légère (fond blanc), flou, bruit gaussien, JPEG q75 -> octets."""
    img = img.convert("RGB").rotate(angle, resample=Image.BICUBIC, expand=False, fillcolor=(255, 255, 255))
    img = img.filter(ImageFilter.GaussianBlur(0.6))
    a = np.asarray(img).astype(np.int16)
    a = a + rng.normal(0, 5, a.shape).astype(np.int16)
    img = Image.fromarray(np.clip(a, 0, 255).astype(np.uint8))
    buf = io.BytesIO(); img.save(buf, "JPEG", quality=75); return buf.getvalue(), img.size


def make_scanned_pdf(ds, out_dir, gt, rows, rng):
    src = fitz.open(os.path.join(ds, "inbound/Dossier_patients_Q3_2026_FICTIF.pdf"))
    pdf_rows = [r for r in rows if r["document"].endswith(".pdf")]
    byp = collections.defaultdict(list)
    for r in pdf_rows:
        byp[int(r["page"])].append(r)
    # meilleure page « courrier » (section letter_*)
    letter = max((p for p in byp if byp[p][0]["section"].startswith("letter")),
                 key=lambda p: sum(r["expected_action"] == "PSEUDONYMIZE" for r in byp[p]))
    pages = [65, 153, letter, 178]           # admin (12 types), facturation (IBAN/cartes), courrier (/Rotate 90), annuaire (dense)
    angles = [1.2, -0.7, 0.9, -1.4]
    doc = fitz.open()
    name = "Rapport_scan_FICTIF.pdf"
    for i, (pno, ang) in enumerate(zip(pages, angles)):
        sp = src[pno - 1]
        pix = sp.get_pixmap(dpi=200)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        w_pt, h_pt = sp.rect.width, sp.rect.height
        if i == 2:
            # page stockée en paysage avec le contenu tourné, /Rotate 90 pour l'affichage (cas fréquent des scanners)
            img = img.rotate(90, expand=True)          # PIL : anti-horaire ; /Rotate 90 (horaire à l'affichage) la redresse
            w_pt, h_pt = h_pt, w_pt
        data, _ = degrade(img, rng, ang)
        page = doc.new_page(width=w_pt, height=h_pt)
        page.insert_image(page.rect, stream=data)
        if i == 2:
            page.set_rotation(90)
        assert page.get_text("words") == [], "la page scannée ne doit pas avoir de couche texte"
        for r in byp[pno]:
            if r["expected_action"] != "PSEUDONYMIZE":
                continue
            gt.n += 1
            gt.rows.append(dict(r, gt_id="GAP%05d" % gt.n, document=name, location_type="image", page=str(i + 1),
                                image_file="", variant="scanned_page(skew=%.1f%s)" % (ang, ",rotate90" if i == 2 else "")))
    doc.save(os.path.join(out_dir, name), deflate=True)
    print("  %s : pages sources %s, skew %s, page 3 /Rotate 90" % (name, pages, angles))
    return pages


# ---------------------------------------------------------------------------------------------------------------------
# A. RTF (DOCX -> LibreOffice -> RTF)
# ---------------------------------------------------------------------------------------------------------------------
def make_letter_docx(ents, gt, path, doc_name):
    from docx import Document
    from docx.shared import Pt
    d = Document()
    P = "P001"
    e = ents[P]
    nm = full_name(ents, P)
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "Patient : %s — Dossier %s" % (nm, e["PATIENT_ID"][0])
    sec.footer.paragraphs[0].text = "%s · %s · Clinique FICTIVE du Léman · Document FICTIF" % (e["PATIENT_ID"][0], nm)
    gt.add_person(doc_name, "text", P, ["FIRST_NAME", "LAST_NAME", "PATIENT_ID"], section="header", variant="rtf_header")
    gt.add_person(doc_name, "text", P, ["FIRST_NAME", "LAST_NAME", "PATIENT_ID"], section="footer", variant="rtf_footer")
    d.add_heading("Courrier de prise en charge — FICTIF", level=1)
    d.add_paragraph("Lausanne, le 12.03.2026")
    d.add_paragraph("Concerne : %s, né le %s, AVS %s, carte d'assuré %s, domicilié %s, %s." % (
        nm, e["DATE_OF_BIRTH"][0], e["AHV_NUMBER"][0], e["INSURANCE_CARD_NUMBER"][0], e["STREET_ADDRESS"][0], e["POSTAL_CITY"][0]))
    gt.add_person(doc_name, "text", P, ["FIRST_NAME", "LAST_NAME", "DATE_OF_BIRTH", "AHV_NUMBER", "INSURANCE_CARD_NUMBER",
                                        "STREET_ADDRESS", "POSTAL_CITY"], section="body", variant="rtf_body")
    d.add_paragraph("Madame, Monsieur,")
    d.add_paragraph("Nous vous confirmons la prise en charge du patient susmentionné. Il est joignable au %s ou par courriel "
                    "à %s. Le remboursement des frais s'effectue sur le compte %s." % (e["PHONE"][0], e["EMAIL"][0], e["IBAN"][0]))
    gt.add_person(doc_name, "text", P, ["PHONE", "EMAIL", "IBAN"], section="body", variant="rtf_body")
    d.add_paragraph("Référence interne : facture FA-2026-10100 (ne pas modifier). Compte de contrôle à checksum invalide, "
                    "à laisser tel quel : CH53 8555 2479 4873 5989 5.")
    gt.decoy(doc_name, "INVOICE_NUMBER", "FA-2026-10100", "n° de facture — pas du PII")
    gt.decoy(doc_name, "IBAN_INVALID", "CH53 8555 2479 4873 5989 5", "forme d'IBAN à checksum faux — ne doit pas être détecté")
    d.add_paragraph("Patients suivis dans le même protocole :")
    t = d.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ("Nom", "Date de naissance", "IBAN")):
        c.text = h
    for eid in ("P002", "P003", "P004", "P005", "P006"):
        row = t.add_row().cells
        row[0].text = full_name(ents, eid); row[1].text = ents[eid]["DATE_OF_BIRTH"][0]; row[2].text = ents[eid]["IBAN"][0]
        gt.add_person(doc_name, "text", eid, ["FIRST_NAME", "LAST_NAME", "DATE_OF_BIRTH", "IBAN"], section="table", variant="rtf_table")
    d.add_paragraph("")
    d.add_paragraph("Avec nos meilleures salutations,")
    d.add_paragraph("Dr Anne Genoud — médecin-cheffe (FICTIF)")
    for p in d.paragraphs:
        for r in p.runs:
            r.font.size = Pt(11)
    d.save(path)


def soffice_convert(soffice, src, fmt, outdir, filter_name=None):
    prof = tempfile.mkdtemp(prefix="lo_gap_")
    target = fmt if not filter_name else "%s:%s" % (fmt, filter_name)
    cmd = [soffice, "--headless", "--norestore", "-env:UserInstallation=file://%s" % prof, "--convert-to", target, "--outdir", outdir, src]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
    shutil.rmtree(prof, ignore_errors=True)
    out = os.path.join(outdir, os.path.splitext(os.path.basename(src))[0] + "." + fmt)
    if not os.path.exists(out):
        raise RuntimeError("LibreOffice: %s\n%s\n%s" % (" ".join(cmd), r.stdout, r.stderr))
    return out


# ---------------------------------------------------------------------------------------------------------------------
# C. imbriqués : classeur, aperçu, DOCX+OLE, PDF+pièce jointe
# ---------------------------------------------------------------------------------------------------------------------
def make_patient_xlsx(ents, gt, path, host_doc, host_loc, eids):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Patients"
    ws.append(["Nom", "Prénom", "AVS", "IBAN"])
    for i, eid in enumerate(eids, start=2):
        e = ents[eid]
        ws.append([e["LAST_NAME"][0], e["FIRST_NAME"][0], e["AHV_NUMBER"][0], e["IBAN"][0]])
        for col, t in zip("ABCD", ("LAST_NAME", "FIRST_NAME", "AHV_NUMBER", "IBAN")):
            gt.add(host_doc, host_loc, t, e[t][0], eid, sheet="Patients", cell="%s%d" % (col, i), variant="embedded_xlsx")
    ws.append(["Contrôle", "—", "756.1234.5678.90", "CH53 8555 2479 4873 5989 5"])   # leurres (checksums faux)
    gt.decoy(host_doc, "AHV_INVALID", "756.1234.5678.90", "AVS à checksum faux dans le classeur imbriqué", sheet="Patients", cell="C%d" % (len(eids) + 2))
    gt.decoy(host_doc, "IBAN_INVALID", "CH53 8555 2479 4873 5989 5", "IBAN à checksum faux dans le classeur imbriqué", sheet="Patients", cell="D%d" % (len(eids) + 2))
    for col, w in zip("ABCD", (18, 16, 20, 30)):
        ws.column_dimensions[col].width = w
    wb.save(path)


def render_table_preview(ents, eids, path):
    """Aperçu façon Word de l'objet Excel : les valeurs en clair dans une image (word/media) — c'est le piège."""
    rows = [("Nom", "Prénom", "AVS", "IBAN")] + [(ents[e]["LAST_NAME"][0], ents[e]["FIRST_NAME"][0], ents[e]["AHV_NUMBER"][0], ents[e]["IBAN"][0]) for e in eids]
    f = font(22); fb = font(22)
    colw = [220, 200, 260, 380]; rh = 40
    W, H = sum(colw) + 20, rh * len(rows) + 20
    img = Image.new("RGB", (W, H), (255, 255, 255)); dr = ImageDraw.Draw(img)
    y = 10
    for ri, row in enumerate(rows):
        x = 10
        for ci, v in enumerate(row):
            dr.rectangle([x, y, x + colw[ci], y + rh], outline=(160, 160, 160), fill=(230, 235, 245) if ri == 0 else None)
            dr.text((x + 8, y + 8), str(v), font=fb if ri == 0 else f, fill=(20, 20, 20))
            x += colw[ci]
        y += rh
    img.save(path)
    return W, H


OLE_RUN = ('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
           'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
           'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">'
           '<w:object w:dxaOrig="{dxa}" w:dyaOrig="{dya}">'
           '<v:shape id="_x0000_i1025" type="#_x0000_t75" style="width:{wpt}pt;height:{hpt}pt" o:ole="">'
           '<v:imagedata r:id="{rimg}" o:title=""/></v:shape>'
           '<o:OLEObject Type="Embed" ProgID="Excel.Sheet.12" ShapeID="_x0000_i1025" DrawAspect="Content" ObjectID="_1790000001" r:id="{remb}"/>'
           '</w:object></w:r>')


def make_embedded_docx(ents, gt, out_dir, tmp):
    from docx import Document
    name = "Contrat_imbrique_FICTIF.docx"
    eids = ["P%03d" % i for i in range(1, 11)]
    P = "P001"; e = ents[P]; nm = full_name(ents, P)
    d = Document()
    d.add_heading("Contrat de participation — étude clinique FICTIVE", level=1)
    d.add_paragraph("Participant principal : %s, né le %s, AVS %s, %s, %s." % (nm, e["DATE_OF_BIRTH"][0], e["AHV_NUMBER"][0],
                                                                              e["STREET_ADDRESS"][0], e["POSTAL_CITY"][0]))
    gt.add_person(name, "paragraph", P, ["FIRST_NAME", "LAST_NAME", "DATE_OF_BIRTH", "AHV_NUMBER", "STREET_ADDRESS", "POSTAL_CITY"],
                  paragraph="1", section="body")
    d.add_paragraph("Le classeur ci-dessous (objet Excel imbriqué) liste les patients inclus dans le protocole :")
    d.add_paragraph("[[OLE_PLACEHOLDER]]")
    d.add_paragraph("Fait à Lausanne le 12.03.2026. Référence contractuelle CT-2026-0042 (ne pas modifier).")
    gt.decoy(name, "CONTRACT_REF", "CT-2026-0042", "référence de contrat — pas du PII")
    d.add_paragraph("Signé : %s" % nm)
    gt.add_person(name, "paragraph", P, ["FIRST_NAME", "LAST_NAME"], paragraph="5", section="body")
    base = os.path.join(tmp, "contrat_base.docx"); d.save(base)

    xlsx = os.path.join(tmp, "Microsoft_Excel_Worksheet1.xlsx")
    make_patient_xlsx(ents, gt, xlsx, name, "embedded", eids)
    prev = os.path.join(tmp, "image_ole_preview.png")
    W, H = render_table_preview(ents, eids, prev)
    for eid in eids:
        gt.add_person(name, "image", eid, ["LAST_NAME", "AHV_NUMBER", "IBAN"], image_file="word/media/image1.png",
                      section="ole_preview", variant="image_ole_preview")

    RIMG, REMB = "rIdGapImg1", "rIdGapEmb1"
    with zipfile.ZipFile(base) as zin, zipfile.ZipFile(os.path.join(out_dir, name), "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                s = data.decode("utf-8")
                # remplacer le run du placeholder par le run <w:object>
                import re
                pat = re.compile(r"<w:r>(?:(?!</w:r>).)*?\[\[OLE_PLACEHOLDER\]\](?:(?!</w:r>).)*?</w:r>", re.S)
                assert pat.search(s), "placeholder introuvable"
                wpt, hpt = 450.0, round(450.0 * H / W, 1)
                s = pat.sub(OLE_RUN.format(dxa=int(wpt * 20), dya=int(hpt * 20), wpt=wpt, hpt=hpt, rimg=RIMG, remb=REMB), s, count=1)
                data = s.encode("utf-8")
            elif item.filename == "word/_rels/document.xml.rels":
                s = data.decode("utf-8")
                rels = ('<Relationship Id="%s" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/package" Target="embeddings/Microsoft_Excel_Worksheet1.xlsx"/>'
                        '<Relationship Id="%s" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image1.png"/>') % (REMB, RIMG)
                data = s.replace("</Relationships>", rels + "</Relationships>").encode("utf-8")
            elif item.filename == "[Content_Types].xml":
                s = data.decode("utf-8")
                add = ""
                if 'Extension="xlsx"' not in s:
                    add += '<Default Extension="xlsx" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"/>'
                if 'Extension="png"' not in s:
                    add += '<Default Extension="png" ContentType="image/png"/>'
                data = s.replace("<Override", add + "<Override", 1).encode("utf-8")
            zout.writestr(item, data)
        zout.write(xlsx, "word/embeddings/Microsoft_Excel_Worksheet1.xlsx")
        zout.write(prev, "word/media/image1.png")
    print("  %s : XLSX imbriqué (word/embeddings) + aperçu PNG en clair (word/media/image1.png)" % name)
    return xlsx, eids


def make_attached_pdf(ents, gt, out_dir, xlsx_src, tmp):
    name = "Annexe_jointe_FICTIF.pdf"
    eids = ["P%03d" % i for i in range(1, 11)]
    P = "P001"; e = ents[P]; nm = full_name(ents, P)
    doc = fitz.open(); page = doc.new_page(width=595, height=842)
    y = 80
    page.insert_text((60, y), "Annexe — Liste des patients du protocole (FICTIF)", fontsize=16, fontname="hebo"); y += 40
    page.insert_text((60, y), "Le classeur complet est joint à ce document (pièce jointe PDF : Liste_patients_FICTIF.xlsx).", fontsize=11); y += 24
    page.insert_text((60, y), "Investigateur de contact : %s — dossier %s — %s" % (nm, e["PATIENT_ID"][0], e["EMAIL"][0]), fontsize=11); y += 24
    page.insert_text((60, y), "Référence CT-2026-0042 · Document FICTIF", fontsize=10, color=(0.4, 0.4, 0.4))
    gt.add_person(name, "text", P, ["FIRST_NAME", "LAST_NAME", "PATIENT_ID", "EMAIL"], page="1", section="body")
    gt.decoy(name, "CONTRACT_REF", "CT-2026-0042", "référence de contrat — pas du PII", page="1")
    xlsx = os.path.join(tmp, "Liste_patients_FICTIF.xlsx")
    make_patient_xlsx(ents, gt, xlsx, name, "embedded", eids)
    doc.embfile_add("Liste_patients_FICTIF.xlsx", open(xlsx, "rb").read(), filename="Liste_patients_FICTIF.xlsx",
                    desc="Classeur des patients (FICTIF)")
    doc.save(os.path.join(out_dir, name), deflate=True)
    print("  %s : 1 page + pièce jointe XLSX (embfile)" % name)


# ---------------------------------------------------------------------------------------------------------------------
# D. signatures
# ---------------------------------------------------------------------------------------------------------------------
def signature_curves(rng, n_seg=7):
    """Paraphe synthétique : suite de Béziers cubiques dans la boîte unité [0,1]x[0,1] + un trait de soulignement."""
    segs = []
    x, y = 0.02, 0.55
    step = 0.9 / n_seg
    for i in range(n_seg):
        nx = x + step
        ny = min(0.95, max(0.05, 0.5 + rng.uniform(-0.35, 0.35)))
        c1 = (x + step * rng.uniform(0.1, 0.5), min(1.0, max(0.0, y + rng.uniform(-0.9, 0.9))))
        c2 = (x + step * rng.uniform(0.5, 0.9), min(1.0, max(0.0, ny + rng.uniform(-0.9, 0.9))))
        segs.append(((x, y), c1, c2, (nx, ny)))
        x, y = nx, ny
    # trait final (paraphe souligné)
    segs.append(((0.05, 0.85), (0.4, 0.98), (0.6, 0.72), (0.95, 0.9)))
    return segs


def bezier_pts(p0, p1, p2, p3, n=40):
    return [((1 - t) ** 3 * p0[0] + 3 * (1 - t) ** 2 * t * p1[0] + 3 * (1 - t) * t ** 2 * p2[0] + t ** 3 * p3[0],
             (1 - t) ** 3 * p0[1] + 3 * (1 - t) ** 2 * t * p1[1] + 3 * (1 - t) * t ** 2 * p2[1] + t ** 3 * p3[1]) for t in (i / n for i in range(n + 1))]


def raster_signature(segs, w, h, stroke=2.5, ink=(20, 25, 90)):
    """Image RGB fond blanc w×h avec le paraphe (anti-aliasé via rendu 4×)."""
    S = 4
    img = Image.new("RGB", (w * S, h * S), (255, 255, 255)); dr = ImageDraw.Draw(img)
    for p0, c1, c2, p3 in segs:
        pts = [(x * w * S, y * h * S) for x, y in bezier_pts(p0, c1, c2, p3)]
        dr.line(pts, fill=ink, width=int(stroke * S), joint="curve")
    return img.resize((w, h), Image.LANCZOS)


def form_page(doc, ents, gt, name, pno, rng, vector):
    """Page A4 : formulaire avec valeurs tapées, libellé « Signature : », paraphe (raster ou vectoriel), nom tapé dessous,
    bloc « Digitally signed by », un cadre et une ligne de séparation ailleurs (leurres graphiques)."""
    P = "P001"; e = ents[P]; nm = full_name(ents, P)
    page = doc.new_page(width=595, height=842)
    page.insert_text((60, 70), "Formulaire de consentement — étude FICTIVE (page %d)" % pno, fontsize=15, fontname="hebo")
    y = 110
    for lab, val in (("Nom, prénom", nm), ("Date de naissance", e["DATE_OF_BIRTH"][0]), ("N° AVS", e["AHV_NUMBER"][0]),
                     ("N° patient", e["PATIENT_ID"][0]), ("Téléphone", e["PHONE"][0])):
        page.insert_text((60, y), lab + " :", fontsize=11); page.insert_text((220, y), val, fontsize=11); y += 22
    gt.add_person(name, "text", P, ["FIRST_NAME", "LAST_NAME", "DATE_OF_BIRTH", "AHV_NUMBER", "PATIENT_ID", "PHONE"], page=str(pno), section="fields")
    # leurre 1 : cadre (encadré d'information) — NE DOIT PAS être recouvert
    frame = fitz.Rect(60, 250, 535, 330)
    page.draw_rect(frame, color=(0.1, 0.1, 0.1), width=1.2)
    page.insert_textbox(fitz.Rect(70, 258, 525, 325), "Information au participant : ce document est entièrement fictif. Il sert à valider "
                        "le masquage des signatures manuscrites sans altérer les cadres et les filets de mise en page.", fontsize=10)
    gt.sig(name, pno, tuple(frame), "decoy_frame", "pt")
    # leurre 2 : ligne de séparation
    page.draw_line((60, 360), (535, 360), color=(0.1, 0.1, 0.1), width=1.0)
    gt.sig(name, pno, (60, 358, 535, 362), "decoy_line", "pt")
    page.insert_text((60, 385), "Le participant confirme avoir lu et compris l'information ci-dessus.", fontsize=10)
    # signature
    page.insert_text((60, 620), "Signature :", fontsize=11, fontname="hebo")
    box = fitz.Rect(150, 585, 360, 645)
    segs = signature_curves(rng)
    if vector:
        for p0, c1, c2, p3 in segs:
            f = lambda p: fitz.Point(box.x0 + p[0] * box.width, box.y0 + p[1] * box.height)
            page.draw_bezier(f(p0), f(c1), f(c2), f(p3), color=(0.08, 0.1, 0.35), width=1.8)
    else:
        img = raster_signature(segs, int(box.width * 200 / 72), int(box.height * 200 / 72), stroke=2.5 * 200 / 72)
        buf = io.BytesIO(); img.save(buf, "PNG")
        page.insert_image(box, stream=buf.getvalue())
    gt.sig(name, pno, tuple(box), "signature_vector" if vector else "signature_raster", "pt")
    page.insert_text((150, 662), nm, fontsize=10)                      # nom tapé sous le paraphe
    page.insert_text((150, 676), "Digitally signed by %s, 12.03.2026" % nm, fontsize=9, color=(0.3, 0.3, 0.3))
    gt.add_person(name, "text", P, ["FIRST_NAME", "LAST_NAME"], page=str(pno), section="typed_name", variant="typed_under_signature")
    gt.add_person(name, "text", P, ["FIRST_NAME", "LAST_NAME"], page=str(pno), section="digital_signature", variant="digitally_signed_by")
    page.insert_text((60, 800), "Document FICTIF — données 100 % synthétiques", fontsize=8, color=(0.5, 0.5, 0.5))
    return page, box, frame


def make_signed_form(ents, gt, out_dir, rng):
    name = "Formulaire_signe_FICTIF.pdf"
    doc = fitz.open()
    p1, box1, frame1 = form_page(doc, ents, gt, name, 1, rng, vector=False)
    form_page(doc, ents, gt, name, 2, rng, vector=True)
    doc.save(os.path.join(out_dir, name), deflate=True)
    # variante image : rendu 200 dpi de la page 1 (boîtes en pixels)
    png = "formulaire_signe_FICTIF.png"
    pix = doc[0].get_pixmap(dpi=200)
    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    img.save(os.path.join(out_dir, png))
    k = 200 / 72
    P = "P001"
    gt.add_person(png, "image", P, ["FIRST_NAME", "LAST_NAME", "DATE_OF_BIRTH", "AHV_NUMBER", "PATIENT_ID", "PHONE"], image_file=png, section="fields", variant="image_printed")
    for s in [s for s in gt.sigs if s["file"] == name and s["page"] == 1]:
        gt.sigs.append(dict(s, file=png, page="", x0=round(s["x0"] * k), y0=round(s["y0"] * k), x1=round(s["x1"] * k), y1=round(s["y1"] * k), unit="px"))
    print("  %s (p.1 paraphe raster, p.2 paraphe vectoriel) + %s (200 dpi)" % (name, png))



# ---------------------------------------------------------------------------------------------------------------------
# Niveau 2 : tests dédiés pour ce qui n'était validé que par lecture du code
#   #3 cohérence texte / tableau / en-tête-pied / image  -> Coherence_FICTIF.docx
#   #4 PII uniquement en en-tête / pied, 4 pages          -> Pied_de_page_FICTIF.docx + Pied_de_page_FICTIF.pdf
#   #1 templates variables                                 -> Template_B_FICTIF.docx (fiche clé/valeur, MAJUSCULES, date DE,
#                                                             IBAN sans espaces) + Template_C_FICTIF.pdf (rapport labo 2 colonnes)
# ---------------------------------------------------------------------------------------------------------------------
def _mapping_has(ds, value):
    with open(os.path.join(ds, "ground_truth/mapping_by_value.csv"), encoding="utf-8") as f:
        return any(r["original_value"] == value for r in csv.DictReader(f))


def _gt_variant(gt, document, loc, dtype, value, rep, eid, **kw):
    """Ligne GT avec une VALEUR VARIANTE (majuscules, sans espaces, date longue) et son pseudonyme attendu."""
    gt.n += 1
    gt.rows.append(dict(gt_id="GAP%05d" % gt.n, document=document, location_type=loc, page=kw.get("page", ""), sheet="", cell="",
                        paragraph=kw.get("paragraph", ""), image_file=kw.get("image_file", ""), section=kw.get("section", ""),
                        data_type=dtype, value=value, entity_id=eid, entity_role="patient", variant=kw.get("variant", "text"),
                        language=kw.get("language", "fr"), expected_action="PSEUDONYMIZE", expected_replacement=rep,
                        strict_replacement=rep, securiti_data_element_hint=HINT.get(dtype, dtype), native_or_custom="natif"))


def make_level2(ds, ents, gt, out_dir, tmp):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt, Inches
    P = "P001"; e = ents[P]; nm = full_name(ents, P)

    # ---- #3 Coherence_FICTIF.docx : la même personne sur 5 surfaces --------------------------------------------------
    name = "Coherence_FICTIF.docx"
    d = Document()
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "Dossier %s — %s" % (nm, e["PATIENT_ID"][0])
    sec.footer.paragraphs[0].text = "%s · AVS %s · Document FICTIF" % (nm, e["AHV_NUMBER"][0])
    d.add_heading("Test de cohérence multi-surfaces (FICTIF)", level=1)
    d.add_paragraph("Le patient %s (AVS %s) est remboursé sur le compte %s." % (nm, e["AHV_NUMBER"][0], e["IBAN"][0]))
    t = d.add_table(rows=1, cols=2); t.style = "Table Grid"
    t.rows[0].cells[0].text = "Champ"; t.rows[0].cells[1].text = "Valeur"
    for k, v in (("Nom, prénom", nm), ("AVS", e["AHV_NUMBER"][0]), ("IBAN", e["IBAN"][0]), ("N° patient", e["PATIENT_ID"][0])):
        row = t.add_row().cells; row[0].text = k; row[1].text = v
    d.add_paragraph("Capture d'écran du système source :")
    img = Image.new("RGB", (1400, 260), (255, 255, 255)); dr = ImageDraw.Draw(img); f = font(34)
    for i, line in enumerate(("Patient : %s" % nm, "AVS : %s" % e["AHV_NUMBER"][0], "IBAN : %s" % e["IBAN"][0])):
        dr.text((30, 25 + i * 75), line, font=f, fill=(20, 20, 20))
    pimg = os.path.join(tmp, "coherence_capture.png"); img.save(pimg)
    d.add_picture(pimg, width=Inches(6))
    d.save(os.path.join(out_dir, name))
    for section, loc in (("header", "text"), ("footer", "text"), ("paragraph", "text"), ("table", "text")):
        types = {"header": ["FIRST_NAME", "LAST_NAME", "PATIENT_ID"], "footer": ["FIRST_NAME", "LAST_NAME", "AHV_NUMBER"],
                 "paragraph": ["FIRST_NAME", "LAST_NAME", "AHV_NUMBER", "IBAN"], "table": ["FIRST_NAME", "LAST_NAME", "AHV_NUMBER", "IBAN", "PATIENT_ID"]}[section]
        gt.add_person(name, loc, P, types, section=section, variant="coherence_" + section)
    gt.add_person(name, "image", P, ["FIRST_NAME", "LAST_NAME", "AHV_NUMBER", "IBAN"], image_file="word/media/image1.png", section="image", variant="coherence_image")
    print("  %s : même personne en en-tête, pied, paragraphe, tableau et image" % name)

    # ---- #4 Pied_de_page_FICTIF.docx : PII UNIQUEMENT en en-tête/pied, 4 pages ------------------------------------------
    name = "Pied_de_page_FICTIF.docx"
    d = Document(); sec = d.sections[0]
    sec.header.paragraphs[0].text = "Confidentiel — %s" % nm
    sec.footer.paragraphs[0].text = "%s · %s · Document FICTIF" % (e["PATIENT_ID"][0], nm)
    for pno in range(1, 5):
        d.add_heading("Protocole clinique FICTIF — section %d" % pno, level=1)
        for _ in range(6):
            d.add_paragraph("Texte de protocole sans donnée personnelle. Posologie standard, suivi hebdomadaire, "
                            "critères d'inclusion et d'exclusion décrits en annexe. Référence interne FA-2026-10100.")
        if pno < 4:
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.save(os.path.join(out_dir, name))
    gt.add_person(name, "text", P, ["FIRST_NAME", "LAST_NAME"], section="header", variant="header_only")
    gt.add_person(name, "text", P, ["FIRST_NAME", "LAST_NAME", "PATIENT_ID"], section="footer", variant="footer_only")
    gt.decoy(name, "INVOICE_NUMBER", "FA-2026-10100", "n° de facture dans le corps — pas du PII")
    print("  %s : 4 pages, PII seulement en en-tête/pied" % name)

    # ---- #4 Pied_de_page_FICTIF.pdf : idem en PDF natif, pied répété sur chaque page ------------------------------------
    name = "Pied_de_page_FICTIF.pdf"
    doc = fitz.open()
    for pno in range(1, 5):
        page = doc.new_page(width=595, height=842)
        page.insert_text((60, 40), "AVS %s — Confidentiel" % e["AHV_NUMBER"][0], fontsize=9, color=(0.3, 0.3, 0.3))
        page.insert_text((60, 90), "Protocole clinique FICTIF — section %d" % pno, fontsize=14, fontname="hebo")
        page.insert_textbox(fitz.Rect(60, 110, 535, 700), ("Texte de protocole sans donnée personnelle. Posologie standard, suivi "
                            "hebdomadaire, critères d'inclusion et d'exclusion décrits en annexe. Référence interne FA-2026-10100. ") * 6, fontsize=10)
        page.insert_text((60, 810), "Patient %s — %s — page %d/4 — Document FICTIF" % (nm, e["PATIENT_ID"][0], pno), fontsize=8, color=(0.4, 0.4, 0.4))
        gt.add_person(name, "text", P, ["AHV_NUMBER"], page=str(pno), section="header", variant="header_only")
        gt.add_person(name, "text", P, ["FIRST_NAME", "LAST_NAME", "PATIENT_ID"], page=str(pno), section="footer", variant="footer_only")
    doc.save(os.path.join(out_dir, name), deflate=True)
    gt.decoy(name, "INVOICE_NUMBER", "FA-2026-10100", "n° de facture dans le corps — pas du PII")
    print("  %s : 4 pages, PII seulement en en-tête/pied" % name)

    # ---- #1 Template_B_FICTIF.docx : fiche clé/valeur paysage, MAJUSCULES, date longue DE, IBAN sans espaces -------------
    name = "Template_B_FICTIF.docx"
    last_up, last_up_rep = e["LAST_NAME"][0].upper(), e["LAST_NAME"][1].upper()
    dob_de, dob_de_rep = "14. März 1962", "5. November 1962"
    iban_ns, iban_ns_rep = e["IBAN"][0].replace(" ", ""), e["IBAN"][1].replace(" ", "")
    assert _mapping_has(ds, last_up) and _mapping_has(ds, dob_de), "variantes attendues dans mapping_by_value.csv"
    d = Document(); sec = d.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = sec.page_height, sec.page_width
    d.add_heading("PATIENTENSTAMMBLATT (FIKTIV) — Vorlage B", level=1)
    t = d.add_table(rows=0, cols=4); t.style = "Light Grid Accent 1"
    pairs = [("Name", last_up), ("Vorname", e["FIRST_NAME"][0]), ("Geburtsdatum", dob_de), ("Patienten-Nr.", e["PATIENT_ID"][0]),
             ("IBAN", iban_ns), ("AHV", e["AHV_NUMBER"][0]), ("Telefon", e["PHONE"][0]), ("E-Mail", e["EMAIL"][0]),
             ("Adresse", e["STREET_ADDRESS"][0]), ("PLZ/Ort", e["POSTAL_CITY"][0]), ("Rechnung", "FA-2026-10100"), ("Studie", "CT-2026-0042")]
    for i in range(0, len(pairs), 2):
        row = t.add_row().cells
        row[0].text, row[1].text = pairs[i]; row[2].text, row[3].text = pairs[i + 1]
        for c in (row[0], row[2]):
            for r_ in c.paragraphs[0].runs: r_.bold = True
    d.add_paragraph("")
    d.add_paragraph("Bemerkung: Herr %s %s wurde am %s geboren; Rückerstattung auf %s." % (last_up, e["FIRST_NAME"][0], dob_de, iban_ns))
    d.save(os.path.join(out_dir, name))
    _gt_variant(gt, name, "text", "LAST_NAME", last_up, last_up_rep, P, section="table", variant="template_B_uppercase", language="de")
    _gt_variant(gt, name, "text", "DATE_OF_BIRTH", dob_de, dob_de_rep, P, section="table", variant="template_B_date_long_de", language="de")
    _gt_variant(gt, name, "text", "IBAN", iban_ns, iban_ns_rep, P, section="table", variant="template_B_iban_nospace", language="de")
    gt.add_person(name, "text", P, ["FIRST_NAME", "PATIENT_ID", "AHV_NUMBER", "PHONE", "EMAIL", "STREET_ADDRESS", "POSTAL_CITY"], section="table", variant="template_B", language="de")
    gt.decoy(name, "INVOICE_NUMBER", "FA-2026-10100", "n° de facture — pas du PII"); gt.decoy(name, "CONTRACT_REF", "CT-2026-0042", "référence étude — pas du PII")
    print("  %s : fiche paysage clé/valeur, MAJUSCULES, date longue DE, IBAN sans espaces" % name)

    # ---- #1 Template_C_FICTIF.pdf : rapport de laboratoire deux colonnes, corps 8 pt, ordre différent --------------------
    name = "Template_C_FICTIF.pdf"
    doc = fitz.open(); page = doc.new_page(width=595, height=842)
    page.insert_text((40, 50), "LABORATOIRE CENTRAL (FICTIF) — Rapport d'analyses — Modèle C", fontsize=12, fontname="hebo")
    page.draw_line((40, 58), (555, 58), width=0.8)
    left = ("Patient : %s %s\nNé le : %s\nN° patient : %s\nAVS : %s\nCourriel : %s\nTél. : %s" %
            (last_up, e["FIRST_NAME"][0], dob_de, e["PATIENT_ID"][0], e["AHV_NUMBER"][0], e["EMAIL"][0], e["PHONE"][0]))
    right = ("Prescripteur : Dr Anne Genoud (FICTIF)\nPrélèvement : 09.09.2026 07:45\nRéf. dossier : LAB-2026-00917\n"
             "Facturation : IBAN %s\nRemarque : à jeun, hémolyse absente" % iban_ns)
    page.insert_textbox(fitz.Rect(40, 70, 300, 200), left, fontsize=8, fontname="helv")
    page.insert_textbox(fitz.Rect(310, 70, 555, 200), right, fontsize=8, fontname="helv")
    y = 220
    page.insert_text((40, y), "Analyte", fontsize=8, fontname="hebo"); page.insert_text((250, y), "Résultat", fontsize=8, fontname="hebo")
    page.insert_text((330, y), "Unité", fontsize=8, fontname="hebo"); page.insert_text((420, y), "Référence", fontsize=8, fontname="hebo")
    for i, (an, res, un, ref) in enumerate((("Glucose", "5.6", "mmol/L", "3.9-5.8"), ("HbA1c", "6.1", "%", "< 5.7"), ("Créatinine", "88", "µmol/L", "62–106"),
                                            ("Potassium", "4.2", "mmol/L", "3.5–5.1"), ("CRP", "3.0", "mg/L", "< 5"), ("Hémoglobine", "14.1", "g/dL", "13.5–17.5"))):
        yy = y + 14 * (i + 1)
        for x, v in ((40, an), (250, res), (330, un), (420, ref)):
            page.insert_text((x, yy), v, fontsize=8)
    page.insert_text((40, 820), "Document FICTIF — jeu de test — aucune donnée réelle", fontsize=7, color=(0.5, 0.5, 0.5))
    doc.save(os.path.join(out_dir, name), deflate=True)
    _gt_variant(gt, name, "text", "LAST_NAME", last_up, last_up_rep, P, page="1", section="body", variant="template_C_uppercase")
    _gt_variant(gt, name, "text", "DATE_OF_BIRTH", dob_de, dob_de_rep, P, page="1", section="body", variant="template_C_date_long_de")
    _gt_variant(gt, name, "text", "IBAN", iban_ns, iban_ns_rep, P, page="1", section="body", variant="template_C_iban_nospace")
    gt.add_person(name, "text", P, ["FIRST_NAME", "PATIENT_ID", "AHV_NUMBER", "EMAIL", "PHONE"], page="1", section="body", variant="template_C")
    gt.decoy(name, "LAB_REF", "LAB-2026-00917", "référence de dossier labo — pas du PII", page="1")
    gt.decoy(name, "LAB_VALUE", "3.9-5.8", "intervalle de référence — pas du PII", page="1")
    print("  %s : rapport labo 2 colonnes, 8 pt, ordre différent" % name)


# ---------------------------------------------------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ds", required=True)
    ap.add_argument("--soffice", default=os.environ.get("SOFFICE", "/Applications/LibreOffice.app/Contents/MacOS/soffice"))
    a = ap.parse_args()
    ds = a.ds
    out_dir = os.path.join(ds, "inbound_gap"); gt_dir = os.path.join(ds, "ground_truth")
    os.makedirs(out_dir, exist_ok=True)
    rng = random.Random(SEED); nrng = np.random.default_rng(SEED)
    rows, ents = load_entities(ds)
    gt = GT(ents)
    tmp = tempfile.mkdtemp(prefix="gap_src_")
    print("Génération dans", out_dir)
    make_scanned_pdf(ds, out_dir, gt, rows, nrng)
    # RTF
    letter_docx = os.path.join(tmp, "Lettre_FICTIF.docx")
    make_letter_docx(ents, gt, letter_docx, "Lettre_FICTIF.rtf")
    if os.path.exists(a.soffice):
        rtf = soffice_convert(a.soffice, letter_docx, "rtf", tmp)
        shutil.copy(rtf, os.path.join(out_dir, "Lettre_FICTIF.rtf"))
        print("  Lettre_FICTIF.rtf : DOCX -> RTF via LibreOffice (%d o)" % os.path.getsize(rtf))
    else:
        shutil.copy(letter_docx, os.path.join(out_dir, "_Lettre_FICTIF.docx.TODO_convert_rtf"))
        print("  !! soffice introuvable (%s) : Lettre_FICTIF.rtf NON généré ; DOCX source déposé pour conversion ultérieure" % a.soffice)
    xlsx_src, _ = make_embedded_docx(ents, gt, out_dir, tmp)
    make_attached_pdf(ents, gt, out_dir, xlsx_src, tmp)
    make_signed_form(ents, gt, out_dir, rng)
    make_level2(ds, ents, gt, out_dir, tmp)
    # CSV
    with open(os.path.join(gt_dir, "ground_truth_gap.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=GT_COLS); w.writeheader()
        for r in gt.rows:
            w.writerow({k: r.get(k, "") for k in GT_COLS})
    with open(os.path.join(gt_dir, "signatures_gap.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["file", "page", "x0", "y0", "x1", "y1", "kind", "unit"]); w.writeheader(); w.writerows(gt.sigs)
    with open(os.path.join(gt_dir, "decoys_gap.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["decoy_id", "document", "page", "section", "sheet", "cell", "kind", "value", "expected", "note"])
        w.writeheader(); w.writerows(gt.decoys)
    by_doc = collections.Counter(r["document"] for r in gt.rows)
    print("ground_truth_gap.csv : %d lignes %s" % (len(gt.rows), dict(by_doc)))
    print("signatures_gap.csv : %d boîtes ; decoys_gap.csv : %d leurres" % (len(gt.sigs), len(gt.decoys)))
    shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    main()
